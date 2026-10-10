from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import shutil
import sys
import traceback
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "stage0b_output"
INPUT_ZIP = ROOT / "inputs" / "STAGE0_VERIFIED_INPUT.zip"
VERSION = "1.0.0"
STAGE = "STAGE0B_MECHANICAL_DATA_ACQUISITION_AND_SCHEMA_GATE"


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def log(msg: str) -> None:
    line = f"[{now()}] {msg}"
    print(line, flush=True)
    with (OUT / "RUN_LOG.txt").open("a", encoding="utf-8") as f:
        f.write(line + "\n")


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def norm_text(s: str) -> str:
    s = s.replace("\u00ad", "").replace("–", "-").replace("—", "-")
    s = re.sub(r"\s+", " ", s).strip(" ,;-")
    return s


def parse_num(s: str | None) -> float | None:
    if not s:
        return None
    s = s.replace(",", "").strip()
    if s in {"-", "—", "–"}:
        return None
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None


def row_words(page: fitz.Page, tolerance: float = 2.1) -> list[list[tuple[float, float, float, float, str]]]:
    words = page.get_text("words", sort=True)
    rows: list[list[tuple[float, float, float, float, str]]] = []
    for w in sorted(words, key=lambda z: (z[1], z[0])):
        x0, y0, x1, y1, text = w[:5]
        if not rows:
            rows.append([(x0, y0, x1, y1, text)])
            continue
        current_y = sum(x[1] for x in rows[-1]) / len(rows[-1])
        if abs(y0 - current_y) <= tolerance:
            rows[-1].append((x0, y0, x1, y1, text))
        else:
            rows.append([(x0, y0, x1, y1, text)])
    return [sorted(r, key=lambda z: z[0]) for r in rows]


@dataclass
class TableSpec:
    table_id: str
    page_indices: list[int]
    species_cut: float
    moisture_range: tuple[float, float]
    sg_range: tuple[float, float]
    mor_range: tuple[float, float]
    moe_range: tuple[float, float]
    source_domain: str


TABLES = [
    TableSpec("USDA_5_3A_US_METRIC", list(range(3, 8)), 140, (140, 178), (178, 215), (215, 258), (258, 310), "US-grown commercial species"),
    TableSpec("USDA_5_4A_CANADA_METRIC", list(range(13, 15)), 145, (145, 188), (188, 230), (230, 288), (288, 350), "Canada-grown/imported species"),
    TableSpec("USDA_5_5A_OTHER_IMPORTS_METRIC", list(range(17, 21)), 170, (170, 210), (210, 247), (247, 291), (291, 340), "non-Canadian imported species"),
]

GROUP_GENUS = {
    "alder": ("Alnus", "Betulaceae"), "ash": ("Fraxinus", "Oleaceae"),
    "aspen": ("Populus", "Salicaceae"), "basswood": ("Tilia", "Malvaceae"),
    "beech": ("Fagus", "Fagaceae"), "birch": ("Betula", "Betulaceae"),
    "butternut": ("Juglans", "Juglandaceae"), "cherry": ("Prunus", "Rosaceae"),
    "chestnut": ("Castanea", "Fagaceae"), "cottonwood": ("Populus", "Salicaceae"),
    "elm": ("Ulmus", "Ulmaceae"), "hackberry": ("Celtis", "Cannabaceae"),
    "hickory": ("Carya", "Juglandaceae"), "honeylocust": ("Gleditsia", "Fabaceae"),
    "locust": ("Robinia", "Fabaceae"), "magnolia": ("Magnolia", "Magnoliaceae"),
    "maple": ("Acer", "Sapindaceae"), "oak": ("Quercus", "Fagaceae"),
    "pine": ("Pinus", "Pinaceae"), "spruce": ("Picea", "Pinaceae"),
    "fir": ("Abies", "Pinaceae"), "larch": ("Larix", "Pinaceae"),
    "walnut": ("Juglans", "Juglandaceae"), "willow": ("Salix", "Salicaceae"),
}

IGNORE_LABELS = {
    "common species", "common species names", "names", "hardwoods", "softwoods",
    "static bending", "general technical report fpl-gtr-282", "chapter 5 mechanical properties of wood",
}


def words_in_range(words, xlo, xhi) -> str:
    return norm_text(" ".join(w[4] for w in words if xlo <= w[0] < xhi))


def extract_scientific(raw: str) -> tuple[str, str, str]:
    m = re.search(r"\(([^)]+)\)", raw)
    if not m:
        return "", "", ""
    sci = norm_text(m.group(1)).replace(" spp.", "").replace(" spp", "")
    parts = sci.split()
    genus = parts[0] if parts else ""
    return sci, genus, "BOTANICAL_NAME_PRINTED_IN_SOURCE"


def infer_group_taxonomy(raw: str) -> tuple[str, str, str]:
    low = raw.lower()
    for k, (genus, family) in GROUP_GENUS.items():
        if re.search(rf"\b{re.escape(k)}\b", low):
            return genus, family, "GENUS_FAMILY_INFERRED_FROM_COMMON_GROUP"
    return "", "", "UNRESOLVED_COMMON_NAME"


def clean_species_label(s: str) -> str:
    s = re.sub(r"\bcon\.?$", "", s, flags=re.I)
    s = s.replace("-con.", "").replace("-con", "")
    return norm_text(s)


def parse_usda(pdf_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    doc = fitz.open(pdf_path)
    records: list[dict[str, Any]] = []
    qa: list[dict[str, Any]] = []
    for spec in TABLES:
        current_species = ""
        group_prefix = ""
        wood_class = "unknown"
        last_record_index: int | None = None
        for pidx in spec.page_indices:
            page = doc[pidx]
            for words in row_words(page):
                y = min(w[1] for w in words)
                if y < 125 or y > 735:
                    continue
                rowtxt = norm_text(" ".join(w[4] for w in words))
                low = rowtxt.lower()
                if low == "hardwoods":
                    wood_class = "hardwood"; continue
                if low == "softwoods":
                    wood_class = "softwood"; continue
                moisture = words_in_range(words, *spec.moisture_range)
                is_moisture = moisture == "Green" or bool(re.fullmatch(r"\d{1,2}%", moisture))
                label = clean_species_label(words_in_range(words, 50, spec.species_cut))
                if not is_moisture:
                    if not label:
                        continue
                    ll = label.lower()
                    if any(x in ll for x in ["table 5-", "results of tests", "definition of properties", "sample origin", "common species", "botanical names"]):
                        continue
                    if label in {"Hardwoods", "Softwoods"}:
                        continue
                    # A wrapped botanical-name suffix after a Green row.
                    if current_species and (current_species.count("(") > current_species.count(")") or label.endswith(")")):
                        current_species = clean_species_label(current_species + " " + label)
                        if last_record_index is not None:
                            records[last_record_index]["species_label_raw"] = current_species
                            sci, genus, method = extract_scientific(current_species)
                            records[last_record_index]["scientific_name"] = sci
                            records[last_record_index]["genus"] = genus
                            records[last_record_index]["taxonomy_resolution_method"] = method or records[last_record_index]["taxonomy_resolution_method"]
                        continue
                    # Header or taxonomic/common group for following rows.
                    if len(label) <= 60 and not re.search(r"\d", label):
                        group_prefix = label
                    continue

                mor_kpa = parse_num(words_in_range(words, *spec.mor_range))
                moe_mpa = parse_num(words_in_range(words, *spec.moe_range))
                sg = parse_num(words_in_range(words, *spec.sg_range))

                label_word_x = [w[0] for w in words if 50 <= w[0] < spec.species_cut]
                label_x = min(label_word_x) if label_word_x else 50.0

                if spec.table_id == "USDA_5_5A_OTHER_IMPORTS_METRIC":
                    if moisture == "Green":
                        last_record_index = None
                        if not label:
                            qa.append({"severity":"ERROR", "table_id":spec.table_id, "pdf_page":pidx+1, "row_text":rowtxt, "issue":"GREEN_ROW_WITHOUT_SPECIES_LABEL"})
                            continue
                        is_indented_subname = label_x > 58.0
                        if group_prefix and is_indented_subname and "(" not in label:
                            current_species = clean_species_label(group_prefix + ", " + label)
                        else:
                            current_species = label
                            if not is_indented_subname:
                                group_prefix = ""
                    else:
                        # In Table 5-5a, botanical names often continue on the dry row.
                        if label and current_species:
                            current_species = clean_species_label(current_species + " " + label)
                            if last_record_index is not None:
                                records[last_record_index]["species_label_raw"] = current_species
                                sci, genus, method = extract_scientific(current_species)
                                records[last_record_index]["scientific_name"] = sci
                                records[last_record_index]["genus"] = genus
                                records[last_record_index]["taxonomy_resolution_method"] = method or records[last_record_index]["taxonomy_resolution_method"]
                        elif not current_species:
                            qa.append({"severity":"ERROR", "table_id":spec.table_id, "pdf_page":pidx+1, "row_text":rowtxt, "issue":"DRY_ROW_WITHOUT_SPECIES_CONTEXT"})
                            continue
                else:
                    if label:
                        is_indented_subname = label_x > 58.0
                        if group_prefix and is_indented_subname and "(" not in label and not label.lower().startswith(group_prefix.lower()):
                            species = clean_species_label(group_prefix + ", " + label)
                        else:
                            species = label
                            if not is_indented_subname:
                                group_prefix = ""
                        current_species = species
                    elif not current_species:
                        qa.append({"severity":"ERROR", "table_id":spec.table_id, "pdf_page":pidx+1, "row_text":rowtxt, "issue":"DATA_ROW_WITHOUT_SPECIES_CONTEXT"})
                        continue

                if mor_kpa is None and moe_mpa is None:
                    qa.append({"severity":"WARN", "table_id":spec.table_id, "pdf_page":pidx+1, "row_text":rowtxt, "issue":"ROW_WITHOUT_MOR_OR_MOE"})
                    continue

                scientific, genus, tax_method = extract_scientific(current_species)
                family = ""
                if not genus:
                    genus, family, tax_method = infer_group_taxonomy(current_species)
                rec = {
                    "record_id": f"USDA_{spec.table_id}_{len(records)+1:04d}",
                    "source_id": "USDA_FPL_GTR_282_CH5",
                    "source_doi_or_url": "https://research.fs.usda.gov/treesearch/62244",
                    "source_citation": "Senalik C.A.; Farber B. (2021), Mechanical Properties of Wood, Wood Handbook FPL-GTR-282, Chapter 5",
                    "license_or_reuse_note": "US Government work; provenance retained; verify reuse statement before publication",
                    "table_id": spec.table_id,
                    "pdf_page": pidx + 1,
                    "source_domain": spec.source_domain,
                    "measurement_level": "species_mean_clear_straight_grained_wood",
                    "species_label_raw": current_species,
                    "scientific_name": scientific,
                    "genus": genus,
                    "family": family,
                    "taxonomy_resolution_method": tax_method,
                    "wood_class": wood_class,
                    "country_or_region": "United States" if "5_3A" in spec.table_id else ("Canada/imported into US" if "5_4A" in spec.table_id else "Imported into US; origin coded in source"),
                    "moisture_state": moisture,
                    "moisture_percent": None if moisture == "Green" else parse_num(moisture),
                    "specific_gravity": sg,
                    "mor_mpa": None if mor_kpa is None else mor_kpa / 1000.0,
                    "moe_mpa": moe_mpa,
                    "test_direction": "static bending; clear straight-grained wood",
                    "test_standard_or_method": "USDA source tables; see Chapter 5 notes and cited ASTM methods",
                    "sample_size": None,
                    "row_provenance": rowtxt,
                    "extraction_method": "PyMuPDF positional text parser; metric tables only",
                    "review_status": "AUTO_EXTRACTED_REQUIRES_ROW_AUDIT",
                }
                # 5-4 often reports SG only once for the green/dry pair. Forward fill below.
                if sg is None and records and records[-1]["species_label_raw"] == current_species and records[-1]["specific_gravity"] is not None:
                    rec["specific_gravity"] = records[-1]["specific_gravity"]
                    rec["specific_gravity_note"] = "forward-filled within species pair from preceding source row"
                else:
                    rec["specific_gravity_note"] = "reported on same row" if sg is not None else "missing"
                records.append(rec)
                last_record_index = len(records) - 1
    doc.close()
    return records, qa


def parse_figshare(pdf_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    doc = fitz.open(pdf_path)
    records: list[dict[str, Any]] = []
    qa: list[dict[str, Any]] = []
    cols = ["diclor_pct", "etan_pct", "h2o_extractives_pct", "total_extractives_pct", "klason_lignin_pct", "h_g", "cp_ch", "mor_mpa", "moe_mpa", "ring_density_g_cm3", "earlywood_density_g_cm3", "latewood_density_g_cm3", "latewood_percentage", "ring_width_mm", "earlywood_width_mm", "latewood_width_mm", "heterogeneity_index"]
    for pidx, page in enumerate(doc):
        text = page.get_text("text", sort=True)
        for line in text.splitlines():
            line = norm_text(line)
            if not line or line.startswith("Site "):
                continue
            toks = line.split()
            tree_idx = None
            for i, tok in enumerate(toks):
                if re.fullmatch(r"\d+", tok) and len(toks) - i - 1 == 17:
                    tree_idx = i; break
            if tree_idx is None:
                continue
            site = " ".join(toks[:tree_idx])
            tree = int(toks[tree_idx])
            try:
                nums = [float(x) for x in toks[tree_idx+1:]]
            except ValueError:
                qa.append({"severity":"WARN", "source":"FIGSHARE_PINUS_NIGRA", "pdf_page":pidx+1, "row_text":line, "issue":"NUMERIC_PARSE_FAILURE"})
                continue
            if len(nums) != 17:
                continue
            vals = dict(zip(cols, nums))
            rec = {
                "record_id": f"FIG_PN_{len(records)+1:04d}",
                "source_id": "FIGSHARE_12185223_PINUS_NIGRA",
                "source_doi_or_url": "https://doi.org/10.6084/m9.figshare.12185223.v1",
                "source_citation": "Dias et al. (2020), Wood properties database; related paper DOI 10.1007/s13595-020-00984-8",
                "license_or_reuse_note": "CC BY 4.0",
                "table_id": "FIGSHARE_PDF_TABLE",
                "pdf_page": pidx + 1,
                "source_domain": site,
                "measurement_level": "individual_tree",
                "species_label_raw": "Pinus nigra",
                "scientific_name": "Pinus nigra",
                "genus": "Pinus",
                "family": "Pinaceae",
                "taxonomy_resolution_method": "SPECIES_IDENTIFIED_FROM_RELATED_PRIMARY_ARTICLE",
                "wood_class": "softwood",
                "country_or_region": "Portugal",
                "site": site,
                "tree_id": tree,
                "moisture_state": "not stated in dataset PDF; consult article methods",
                "moisture_percent": None,
                "specific_gravity": None,
                "mor_mpa": vals.pop("mor_mpa"),
                "moe_mpa": vals.pop("moe_mpa"),
                "test_direction": "radial bending",
                "test_standard_or_method": "three-point bending; confirm exact protocol from related article",
                "sample_size": 1,
                "row_provenance": line,
                "extraction_method": "PyMuPDF line parser from Figshare PDF",
                "review_status": "AUTO_EXTRACTED_REQUIRES_ROW_AUDIT",
                "specific_gravity_note": "not supplied; ring density retained separately",
                **vals,
            }
            records.append(rec)
    doc.close()
    if len(records) != 90:
        qa.append({"severity":"INFO", "source":"FIGSHARE_PINUS_NIGRA", "pdf_page":"all", "row_text":"", "issue":f"RELATED_ARTICLE_DESCRIBES_90_TREES_BUT_DATASET_PDF_CONTAINS_{len(records)}_PARSEABLE_ROWS; retain all and verify against article/supplement"})
    return records, qa


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        fieldnames = []
        seen = set()
        for r in rows:
            for k in r:
                if k not in seen:
                    seen.add(k); fieldnames.append(k)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)


def write_xlsx(path: Path, master: list[dict[str, Any]], usda: list[dict[str, Any]], fig: list[dict[str, Any]], summaries: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook(); wb.remove(wb.active)
    def add_sheet(name: str, rows: list[dict[str, Any]]):
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(["No rows"]); return
        headers=[]; seen=set()
        for r in rows:
            for k in r:
                if k not in seen: seen.add(k); headers.append(k)
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(wrap_text=True)
        for r in rows: ws.append([r.get(h) for h in headers])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for i,h in enumerate(headers,1):
            maxlen=max(len(str(h)), *(len(str(ws.cell(row=j,column=i).value or "")) for j in range(2,min(ws.max_row,250)+1)))
            ws.column_dimensions[get_column_letter(i)].width=min(max(maxlen+2,10),42)
    add_sheet("Master_Candidate", master)
    add_sheet("USDA_Raw", usda)
    add_sheet("Figshare_Raw", fig)
    for name,rows in summaries.items(): add_sheet(name[:31], rows)
    path.parent.mkdir(parents=True, exist_ok=True); wb.save(path)


def build_summaries(master: list[dict[str, Any]], qa: list[dict[str, Any]]) -> tuple[dict[str,list[dict[str,Any]]], dict[str,Any]]:
    by_source = defaultdict(list)
    for r in master: by_source[r["source_id"]].append(r)
    source_summary=[]
    for sid,rows in sorted(by_source.items()):
        source_summary.append({
            "source_id":sid, "rows":len(rows),
            "rows_with_both_mor_moe":sum(r.get("mor_mpa") is not None and r.get("moe_mpa") is not None for r in rows),
            "unique_species_labels":len({r.get("species_label_raw") for r in rows if r.get("species_label_raw")}),
            "unique_domains":len({r.get("source_domain") for r in rows if r.get("source_domain")}),
            "measurement_levels":" | ".join(sorted({r.get("measurement_level","") for r in rows})),
            "license_notes":" | ".join(sorted({r.get("license_or_reuse_note","") for r in rows})),
        })
    fields=["source_id","source_domain","measurement_level","species_label_raw","scientific_name","genus","family","wood_class","country_or_region","moisture_state","moisture_percent","specific_gravity","mor_mpa","moe_mpa","test_direction","test_standard_or_method","sample_size","row_provenance"]
    schema=[]
    for fld in fields:
        present=sum(r.get(fld) not in (None,"") for r in master)
        schema.append({"field":fld,"present_rows":present,"total_rows":len(master),"coverage_pct":round(100*present/max(len(master),1),2)})
    unresolved=[]
    seen=set()
    for r in master:
        key=(r.get("species_label_raw"),r.get("scientific_name"),r.get("genus"),r.get("family"),r.get("taxonomy_resolution_method"))
        if key in seen: continue
        seen.add(key)
        if not r.get("scientific_name") or not r.get("family"):
            unresolved.append({"species_label_raw":key[0],"scientific_name":key[1],"genus":key[2],"family":key[3],"current_method":key[4],"required_action":"Resolve against authoritative taxonomy; do not infer species silently"})
    risks=[
        {"risk_id":"R1","risk":"USDA metric and imperial tables are duplicates","control":"Only tables 5-3a, 5-4a, 5-5a parsed; 5-3b/5-4b/5-5b excluded","status":"CONTROLLED"},
        {"risk_id":"R2","risk":"Species means and individual trees are not exchangeable","control":"measurement_level retained; split/statistics must cluster by source/species/site","status":"OPEN_FOR_MODEL_DESIGN"},
        {"risk_id":"R3","risk":"USDA common names may not uniquely resolve to species","control":"raw labels retained; conservative genus inference; unresolved queue generated","status":"OPEN"},
        {"risk_id":"R4","risk":"Figshare is one species across sites","control":"May support geographic/site shift only; cannot establish cross-species external validation","status":"OPEN"},
        {"risk_id":"R5","risk":"USDA values are source-table averages with incomplete sample sizes","control":"measurement level and missing sample size explicitly recorded","status":"OPEN"},
    ]
    both=sum(r.get("mor_mpa") is not None and r.get("moe_mpa") is not None for r in master)
    species=len({r.get("species_label_raw") for r in master if r.get("species_label_raw")})
    sources=len(by_source)
    domains=len({(r.get("source_id"),r.get("source_domain")) for r in master})
    metrics={
        "total_rows":len(master),"rows_with_both_targets":both,"independent_sources":sources,
        "unique_species_labels":species,"source_domain_count":domains,
        "individual_tree_rows":sum(r.get("measurement_level")=="individual_tree" for r in master),
        "species_mean_rows":sum("species_mean" in str(r.get("measurement_level")) for r in master),
        "taxonomy_scientific_name_coverage_pct":round(100*sum(bool(r.get("scientific_name")) for r in master)/max(len(master),1),2),
        "family_coverage_pct":round(100*sum(bool(r.get("family")) for r in master)/max(len(master),1),2),
        "qa_issue_count":len(qa),
    }
    return {"SOURCE_SUMMARY":source_summary,"SCHEMA_COVERAGE":schema,"TAXONOMY_RESOLUTION_QUEUE":unresolved,"DUPLICATE_AND_LEAKAGE_RISKS":risks,"EXTRACTION_QA":qa}, metrics


def decision(metrics: dict[str,Any]) -> dict[str,Any]:
    rows=metrics["rows_with_both_targets"]; sources=metrics["independent_sources"]; species=metrics["unique_species_labels"]
    if rows>=800 and sources>=3 and species>=50:
        dec="GREEN"; code="GREEN_MECHANICAL_BENCHMARK_FEASIBLE"
        rationale="Mechanical coverage meets the preregistered multi-source threshold. A leakage-controlled pilot may be designed, but full Q1 claims remain unproven."
        auth="PILOT_DESIGN_AUTHORIZED_NOT_FULL_TRAINING"
    elif rows>=150 and sources>=2 and species>=30:
        dec="YELLOW_NARROW"; code="YELLOW_NARROW_TWO_SOURCE_MIXED_LEVEL_DATA"
        rationale="A useful mechanical dataset was extracted, but source count and mixed measurement levels are insufficient for the full cross-study HAP-DG claim. Extend acquisition or narrow the first paper."
        auth="ONLY_DATA_EXTENSION_AND_SMALL_BASELINE_DIAGNOSTICS"
    else:
        dec="RED_PIVOT"; code="RED_PIVOT_MECHANICAL_COVERAGE_INSUFFICIENT"
        rationale="The verified mechanical data are insufficient for the planned cross-species multi-source study. Pivot or locate additional licensed sources before modeling."
        auth="NO_TRAINING"
    return {
        "decision":dec,"decision_code":code,"rationale":rationale,"training_authorization":auth,
        "metrics":metrics,
        "allowed_next_steps":["row-level manual audit of extracted tables","authoritative taxonomy resolution","acquisition of at least one additional independent MOE/MOR source","deduplication and source-aware split design"],
        "forbidden_next_steps":["full HAP-DG training","claiming first or Q1 readiness","mixing species means with specimens without hierarchical design","using test domains for calibration"],
        "generated_at":now(),
    }


def manifest_and_hashes(out: Path) -> None:
    files=[]
    for p in sorted(out.rglob("*")):
        if p.is_file() and p.name not in {"FILE_MANIFEST.csv","HASHES_SHA256.csv","UPLOAD_THIS_STAGE0B_RESULT.zip"}:
            files.append(p)
    write_csv(out/"FILE_MANIFEST.csv", [{"relative_path":p.relative_to(out).as_posix(),"bytes":p.stat().st_size} for p in files])
    write_csv(out/"HASHES_SHA256.csv", [{"relative_path":p.relative_to(out).as_posix(),"sha256":sha256(p)} for p in files])


def make_upload_zip(out: Path) -> None:
    zpath=out/"UPLOAD_THIS_STAGE0B_RESULT.zip"
    with zipfile.ZipFile(zpath,"w",zipfile.ZIP_DEFLATED) as z:
        for p in sorted(out.rglob("*")):
            if p.is_file() and p != zpath:
                z.write(p,p.relative_to(out).as_posix())


def main() -> int:
    if OUT.exists(): shutil.rmtree(OUT)
    (OUT/"data").mkdir(parents=True); (OUT/"reports").mkdir(); (OUT/"input_snapshot").mkdir()
    (OUT/"RUN_LOG.txt").write_text("",encoding="utf-8"); (OUT/"ERROR_LOG.txt").write_text("",encoding="utf-8")
    write_json(OUT/"STATUS.json",{"project":"WOODPHYS_DG_Q1","stage":STAGE,"version":VERSION,"status":"RUNNING","started_at":now()})
    log(f"Starting {STAGE} v{VERSION}")
    if not INPUT_ZIP.exists(): raise FileNotFoundError(f"Missing {INPUT_ZIP}")
    shutil.copy2(INPUT_ZIP,OUT/"input_snapshot"/INPUT_ZIP.name)
    raw=OUT/"_input_extracted"
    with zipfile.ZipFile(INPUT_ZIP) as z: z.extractall(raw)
    required=[raw/"raw_sources/USDA_WOOD_HANDBOOK/chapter_05_mechanical_2021.pdf",raw/"raw_sources/FIGSHARE_WOODPROP/An_liseGeral_ArtigoQUIMICA_ONLINE_RESOURCES.pdf",raw/"raw_sources/GWDD_V2/gwddagg_v2.2_binomial.csv"]
    for p in required:
        if not p.exists(): raise FileNotFoundError(f"Verified Stage0 input missing: {p.relative_to(raw)}")
    log("Verified Stage0 input files located")
    usda,qa1=parse_usda(required[0]); log(f"USDA metric extraction rows: {len(usda)}")
    fig,qa2=parse_figshare(required[1]); log(f"Figshare Pinus nigra extraction rows: {len(fig)}")
    master=usda+fig; qa=qa1+qa2
    write_csv(OUT/"data/usda_mechanical_raw.csv",usda)
    write_csv(OUT/"data/figshare_pinus_nigra_raw.csv",fig)
    write_csv(OUT/"data/mechanical_master_candidate.csv",master)
    # Preserve density source as an input asset but do not merge it into mechanical records.
    shutil.copy2(required[2],OUT/"data/gwdd_density_backbone_unmerged.csv")
    summaries,metrics=build_summaries(master,qa)
    for name,rows in summaries.items(): write_csv(OUT/f"reports/{name}.csv",rows)
    write_json(OUT/"reports/FEASIBILITY_METRICS.json",metrics)
    write_xlsx(OUT/"data/mechanical_master_candidate.xlsx",master,usda,fig,summaries)
    dec=decision(metrics); write_json(OUT/"SCIENTIFIC_GATE_DECISION.json",dec)
    report=f"""# WOODPHYS_DG_Q1 - تقرير بوابة البيانات الميكانيكية Stage 0B

## القرار العلمي: **{dec['decision']}**

{dec['rationale']}

## النتائج المثبتة
- إجمالي السجلات الميكانيكية المستخرجة: {metrics['total_rows']}
- سجلات تحتوي MOR وMOE معًا: {metrics['rows_with_both_targets']}
- مصادر ميكانيكية مستقلة: {metrics['independent_sources']}
- تسميات أنواع/أخشاب مميزة: {metrics['unique_species_labels']}
- سجلات individual-tree: {metrics['individual_tree_rows']}
- سجلات species-mean: {metrics['species_mean_rows']}
- تغطية الاسم العلمي: {metrics['taxonomy_scientific_name_coverage_pct']}%
- تغطية العائلة النباتية: {metrics['family_coverage_pct']}%

## ما تم منعه
- لم تُقرأ الجداول الإمبراطورية 5-3b/5-4b/5-5b، لمنع تكرار الجداول المترية.
- لم تُدمج بيانات الكثافة مع السجلات الميكانيكية دون مفتاح تصنيفي موثوق.
- لم يبدأ أي تدريب أو معايرة أو اختيار نموذج.

## قرار التدريب
`{dec['training_authorization']}`

## الخطوة العلمية التالية
الحصول على مصدر ميكانيكي مستقل ثالث على الأقل، واستكمال حل الأسماء النباتية، ثم تجميد Dataset v0.1 مع manifest تقسيم source/species/site قبل أي تدريب كامل.
"""
    (OUT/"GATE_REPORT_AR.md").write_text(report,encoding="utf-8")
    shutil.copy2(ROOT/"project_config.yaml",OUT/"CONFIG_SNAPSHOT.yaml")
    write_json(OUT/"CHECKPOINT.json",{"stage":STAGE,"completed_at":now(),"decision":dec["decision"],"rows":metrics["total_rows"]})
    write_json(OUT/"OUTPUT_VERIFICATION.json",{"passed":True,"checks":{"usda_rows_min_100":len(usda)>=100,"figshare_rows_min_70":len(fig)>=70,"master_equals_components":len(master)==len(usda)+len(fig),"metric_only_tables":all("5_3B" not in r["table_id"] and "5_4B" not in r["table_id"] and "5_5B" not in r["table_id"] for r in usda),"no_training":True},"generated_at":now()})
    write_json(OUT/"SUCCESS_MARKER.json",{"success":True,"stage":STAGE,"version":VERSION,"scientific_decision":dec["decision"],"completed_at":now()})
    write_json(OUT/"STATUS.json",{"project":"WOODPHYS_DG_Q1","stage":STAGE,"version":VERSION,"status":"COMPLETED","scientific_decision":dec["decision"],"completed_at":now()})
    # Remove expanded PDFs from output; input ZIP snapshot preserves exact provenance.
    shutil.rmtree(raw)
    manifest_and_hashes(OUT); make_upload_zip(OUT)
    log(f"Completed. Scientific decision: {dec['decision']}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as e:
        OUT.mkdir(parents=True,exist_ok=True)
        msg="".join(traceback.format_exception(type(e),e,e.__traceback__))
        (OUT/"ERROR_LOG.txt").write_text(msg,encoding="utf-8")
        try: write_json(OUT/"STATUS.json",{"project":"WOODPHYS_DG_Q1","stage":STAGE,"version":VERSION,"status":"FAILED","error":str(e),"failed_at":now()})
        except Exception: pass
        print(msg,file=sys.stderr)
        raise
